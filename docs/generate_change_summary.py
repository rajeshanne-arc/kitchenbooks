from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = "docs/kitchenbooks-detailed-change-summary.docx"
BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)
MUTED = RGBColor(89, 89, 89)


def set_font(run, name="Calibri", size=11, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    p.add_run(text)
    return p


def add_section(doc, number, title, why, items):
    h = doc.add_heading(f"{number}. {title}", level=1)
    h.paragraph_format.keep_with_next = True
    p = doc.add_paragraph()
    r = p.add_run("Why this was done: ")
    set_font(r, bold=True, color=DARK)
    p.add_run(why)
    for item in items:
        add_bullet(doc, item)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10

for name, size, color, before, after in [
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, DARK, 8, 4),
]:
    style = styles[name]
    style.font.name = "Calibri"
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = True
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.text = "KitchenBooks  |  Engineering Change Summary"
for run in header.runs:
    set_font(run, size=9, color=MUTED)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run("KitchenBooks")
set_font(run, size=9, color=MUTED)

title = doc.add_paragraph()
title.paragraph_format.space_after = Pt(4)
r = title.add_run("KitchenBooks")
set_font(r, size=26, color=RGBColor(0, 0, 0), bold=True)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(16)
r = subtitle.add_run("Detailed Change Summary and Rationale")
set_font(r, size=15, color=MUTED)

meta = doc.add_paragraph()
meta.paragraph_format.space_after = Pt(12)
for label, value in [
    ("Project: ", "rajeshanne-arc/kitchenbooks\n"),
    ("Deployment: ", "Self-hosted server at kb.etdemo.in\n"),
    ("Status: ", "Technical implementation and isolated hosted acceptance completed; final business approval remains pending"),
]:
    r = meta.add_run(label)
    set_font(r, size=10.5, bold=True)
    r = meta.add_run(value)
    set_font(r, size=10.5, color=MUTED)

note = doc.add_paragraph()
note.paragraph_format.space_after = Pt(12)
r = note.add_run("Important scope note: ")
set_font(r, bold=True, color=DARK)
note.add_run("These are changes made in the working tree. They have not been committed or pushed. Secrets and passwords are intentionally excluded from this report.")

add_section(doc, 1, "Production product specification", "The work needed a durable source of truth so implementation could continue without losing product context.", [
    "Created and updated product-spec.md, feature-status.md, worklog.md, production-release-checklist.md, and production-approval.md.",
    "Recorded the agreed ground rule: continue end-to-end and stop only for a genuinely external dependency or an irreversible live-data approval.",
    "Separated implemented functionality, technical verification, external dependencies, and business sign-off.",
])

add_section(doc, 2, "Self-hosted production deployment", "The shared Supabase project was unsuitable for the application because its database credentials and usage were shared by multiple people.", [
    "Deployed KitchenBooks on the authorized server at 100.64.33.51.",
    "Configured kitchenbooks.service to run the application on 127.0.0.1:3120.",
    "Reused the existing PostgreSQL Docker cluster, but created a separate logical database named kitchenbooks.",
    "Added a local database TLS proxy on 127.0.0.1:55432.",
    "Did not create a second Docker container.",
])

add_section(doc, 3, "Cloudflare routing", "The application needed a public HTTPS address while remaining bound to localhost on the server.", [
    "Configured the Cloudflare Tunnel route kb.etdemo.in to forward to localhost:3120.",
    "Verified that https://kb.etdemo.in serves the application.",
    "Kept the application port private rather than exposing it directly to the internet.",
])

add_section(doc, 4, "Database and migrations", "The new deployment needed a complete, independently verifiable KitchenBooks database rather than relying on incomplete shared tables.", [
    "Applied and verified the complete migration set on the self-hosted kitchenbooks database.",
    "Added structures for tenants, users, memberships, roles, inventory, lots, purchases, sales, accounting, payroll, leave, attachments, approvals, and business days.",
    "Replayed and audited migrations without modifying the shared Supabase project.",
])

add_section(doc, 5, "Tenant isolation and database security", "Restaurant data must remain isolated even if an application route or client request is incorrect.", [
    "Strengthened tenant-aware access throughout the database.",
    "Added or verified row-level security across tenant tables.",
    "Added tenant-aware security-invoker behavior for sensitive views.",
    "Added composite foreign keys containing tenant identifiers.",
    "Made business-date operations tenant-aware.",
    "Added strict tenancy audit scripts and verified them successfully.",
])

add_section(doc, 6, "Users, roles, and permissions", "KitchenBooks has different operational responsibilities, so each user must see and perform only the work appropriate to their role.", [
    "Implemented and tested Owner, Manager, Chef/Kitchen, Store, Cashier/Sales, Staff, and Accountant roles.",
    "Added membership-based tenant access.",
    "Added invitations and password reset support.",
    "Protected owner-only setup and user-management screens.",
    "Added role-specific dashboard routing and unauthorized-access handling.",
])

add_section(doc, 7, "Login and session fixes", "The application could authenticate successfully but then fail on the next page because membership and tenant context were resolved incorrectly.", [
    "Changed membership lookup to use secure database functions instead of directly querying rows hidden by row-level security.",
    "Fixed references to membership fields that were not returned by the database function.",
    "Restored tenant context from the signed session cookie when a business-day query did not already have a tenant context.",
    "Verified the fix against the hosted owner and role pages.",
])

add_section(doc, 8, "Inventory and stock-lot accounting", "Food inventory needs a traceable ledger and must prevent unexplained or impossible stock balances.", [
    "Added append-only stock-lot movements.",
    "Added lot creation on receiving and FEFO-style lot selection.",
    "Prevented over-issuing stock that was never received.",
    "Added lot-level audit history and controlled adjustments.",
    "Added stock valuation and invalid-balance checks.",
])

add_section(doc, 9, "Purchasing workflow", "Purchases should be reviewed and committed atomically so inventory and accounting cannot become partially updated.", [
    "Implemented purchase orders, approvals, vendor quotes, and invoice matching.",
    "Added grouped purchase-batch imports with preview before commit.",
    "Added vendor and bill references.",
    "Made purchase commits atomic.",
    "Connected receiving and purchase records to accounting mappings.",
])

add_section(doc, 10, "Sales and POS reconciliation", "POS sales, preparation demand, and inventory movements are related but should not be conflated automatically.", [
    "Added sales capture and POS statement import support.",
    "Added a preview gate before POS data is committed.",
    "Added reconciliation support.",
    "Added an owner-controlled POS stock policy: no stock movement or stock reconciliation.",
    "Added mapped preparation-demand behavior.",
])

add_section(doc, 11, "Accounting", "Financial records need clear posting rules, traceability, and controlled corrections.", [
    "Added chart of accounts and posting mappings.",
    "Added accounting journals and journal lines.",
    "Added business-day close behavior.",
    "Added accrual and fixed-asset support.",
    "Added register and recurring-accounting screens.",
    "Implemented immutable void-and-refile correction behavior instead of silently overwriting posted records.",
])

add_section(doc, 12, "Payroll and leave", "Payroll and leave balances require structured records and validation rather than manual totals.", [
    "Added payroll workflow support.",
    "Added a formal leave ledger and carry-forward calculations.",
    "Added PF/ESI validation.",
    "Added neutral withholding CSV export.",
])

add_section(doc, 13, "SOPs and localization", "The product is intended for operational teams with different responsibilities and language needs.", [
    "Added role-specific SOP moments across kitchen, store, sales, staff, owner, and accounts workflows.",
    "Added Telugu localization support.",
    "Improved user-facing workflow labels and explanations.",
])

add_section(doc, 14, "Attachment storage", "Bill photos and attachments needed to work without external Blob or R2 credentials.", [
    "Added private local filesystem storage controlled by KB_FILE_STORAGE_DIR.",
    "Restricted storage directory and file permissions.",
    "Rejected path traversal attempts.",
    "Served attachments through authenticated application routes.",
    "Preserved the existing Vercel Blob path so storage can later move to Blob or R2 without redesigning the feature.",
])

add_section(doc, 15, "Petpooja and demo data", "Real Petpooja credentials were unavailable, so the integration needed a safe way to be tested without pretending to access a real provider account.", [
    "Added Petpooja configuration and demo-mode support.",
    "Added isolated seed/probe data for testing provider-related workflows.",
    "Kept real provider credentials as an external dependency rather than inventing them.",
])

add_section(doc, 16, "Public landing page", "Unauthenticated visitors needed a clear explanation of the product instead of being sent immediately to a blank or login-only experience.", [
    "Replaced the root page with a production-style KitchenBooks landing page.",
    "Added sections for purchasing, inventory, kitchen operations, sales, accounting, payroll, owner controls, evidence, imports, SOPs, and language support.",
    "Added role cards for Owner, Manager, Chef, Store, Cashier, and Accountant.",
    "Added How it works, Roles, and Sign in navigation.",
])

add_section(doc, 17, "Public route handling", "The middleware previously redirected the root path before the new landing page could render.", [
    "Marked the root path as public in src/proxy.ts.",
    "Kept protected application routes behind authentication.",
    "Kept authenticated users redirected to their role dashboard.",
])

add_section(doc, 18, "Testing and smoke-test corrections", "Automated checks must match the current product rules and must never write into live business data.", [
    "Added migration, schema, RLS, tenancy, role-matrix, stock-lot, deployment, and preflight audits.",
    "Restricted write-capable smoke tests to an isolated probe tenant.",
    "Updated store tests to use the actual configured receiving sections instead of a stale hardcoded count.",
    "Updated dashboard tests to validate lot-ledger over-issue rejection instead of expecting an impossible negative balance.",
    "Verified purchase-batch, sales, store, and dashboard flows on the hosted probe.",
])

add_section(doc, 19, "Verification completed", "The implementation needed evidence that it builds, starts, serves traffic, and passes the core technical gates.", [
    "TypeScript compilation passed.",
    "Lint passed.",
    "Production build passed.",
    "Dependency audit passed with no high-severity production vulnerabilities.",
    "Migration replay, schema audit, strict tenancy audit, access-matrix audit, and stock-lot audit passed.",
    "Hosted public, login, owner, kitchen, store, sales, staff, and accounts routes were verified.",
    "Cloudflare routing, systemd services, local attachment storage, and hosted probe workflows were verified.",
])

add_section(doc, 20, "What remains outside the implementation", "Some items cannot be completed truthfully without external credentials or an authorized business decision.", [
    "Real Petpooja credentials and provider validation are still required.",
    "WhatsApp credentials and template approval are still required if WhatsApp notifications are part of launch scope.",
    "A real business-day acceptance using owner-approved operational data is still required.",
    "The owner/accountant must provide the final production approval; it was not forged or self-signed.",
    "The current attachment backend is local filesystem storage. R2 can replace it later behind the same interface.",
])

doc.add_heading("Source files and records", level=1)
for path in [
    "docs/product-spec.md",
    "docs/feature-status.md",
    "docs/worklog.md",
    "docs/production-release-checklist.md",
    "docs/production-approval.md",
    "docs/attachments-storage-decision.md",
    "src/app/page.tsx",
    "src/components/TopNav.tsx",
    "src/proxy.ts",
    "src/server/blob.ts",
    "src/server/current-user.ts",
    "src/server/business-day.ts",
]:
    add_bullet(doc, path)

doc.add_page_break()
doc.add_heading("Appendix A — Detailed functional change ledger", level=1)
p = doc.add_paragraph("This appendix expands the workstream summary into the individual user-facing and domain-level functions implemented or wired during the project. It intentionally excludes credentials and secret values.")

functional_groups = [
    ("Identity, access, and restaurant lifecycle", [
        "Self-service password change at /account/security.",
        "Owner-issued one-time password reset flow with post-reset confirmation and first-use guidance.",
        "Global identity plus restaurant-membership model while preserving compatibility with existing app_users records.",
        "Global-login restaurant selection when KB_MEMBERSHIPS=true.",
        "Restaurant switching with signed-session revalidation against active memberships.",
        "Owner controls to add restaurant access, list memberships, retire memberships, and restore memberships.",
        "Database safeguard preventing removal of the last active owner.",
        "Secure invitations for new global accounts: 48-hour expiry, one-use consumption, hashed token storage, database-side username/staff/role validation, and a public /invite acceptance route.",
        "Platform provisioning endpoint that atomically creates a restaurant, baseline settings, departments, and the first owner.",
        "Database synchronization trigger keeping legacy account actions consistent with global identities and memberships.",
    ]),
    ("Purchasing and vendor operations", [
        "Purchase orders with draft, submit-for-approval, pending, approved, refused, and send states.",
        "Approval queue showing amount, requester, reason, and decision controls.",
        "Owner setting for disabling PO approval or applying a rupee approval threshold.",
        "Pending purchase orders are frozen; refusal restores editability; approval unlocks normal vendor sending.",
        "Vendor quotation register with quote headers and lines, dates, and evidence status.",
        "Accepted quotation conversion into a draft purchase order while retaining the quote as immutable evidence.",
        "Purchase invoice matching with matched, partial, exception, and unmatched states, quantity/rate flags, and comparison snapshots.",
        "One-bill purchase CSV import with exact validation, active vendor/item resolution, preview, explicit commit, and delegation to the normal atomic bill transaction.",
        "Historical grouped purchase-batch import keyed by vendor and supplier bill number, with duplicate-reference validation, source bill retention, full-file preview, all-or-nothing commit, and rollback probing.",
        "Vendor-master import with exact headers, field validation, duplicate detection, sequential category-specific codes, preview, and atomic commit.",
        "Item-master import with category, unit, storage location, rate, expiry, duplicate, and existing-name validation, preview, code assignment, and atomic commit.",
    ]),
    ("Inventory, stock, and locations", [
        "Tenant-scoped stock lots with FEFO ordering indexes and append-only signed lot movements.",
        "Legacy aggregate on-hand represented by clearly labelled legacy lots rather than fabricated historical FIFO provenance.",
        "New purchase receipts create provenance-rich lots.",
        "Stock issues allocate earliest-expiry lots inside the issue transaction.",
        "Issue reversals restore the recorded lot allocations.",
        "Vendor returns post negative movements against the source receipt lot and reject over-returns.",
        "Purchase-bill voids reverse all source receipt lots through signed reversal movements.",
        "Location transfers record source and destination movements, preserve total stock/value, and validate source availability.",
        "Unplaced lots remain eligible for FEFO issue and wastage; transfers still require a real source location.",
        "Wastage reversals use their own movement type.",
        "Opening-stock CSV import with BOM/quoted-field support, write-free preview, duplicate-free quantities, active-item resolution, and explicit commit through the normal adjustment transaction.",
        "Owner-configurable stock-adjustment approval path with frozen request lines, server-captured unit cost, approval/refusal queue, and stock write only after approval.",
        "Stock-lot audit and aggregate-to-legacy-lot reconciliation checks.",
    ]),
    ("Recipes, kitchen, production, and SOPs", [
        "Immutable recipe versions with effective dates and history display.",
        "Production entries pinned to the recipe version effective on the production date; reversals retain the same pinned version.",
        "Recipe-line substitutions restricted to active items, positive quantity ratios, and a specific ingredient line.",
        "Substitutions cannot target sub-recipe lines and are included in the next recipe-version snapshot.",
        "Production expected-output snapshots and measured-waste inputs.",
        "Monthly production variance review showing expected, made, waste, and signed variance by recipe.",
        "Manager/owner variance acknowledgement and correction requests requiring a note and preserving the original production record.",
        "Mapped POS preparation-demand surface showing dish portions sold for the current month and linking to recipe cards.",
        "Preparation demand is planning-only and never creates stock issues or ingredient consumption.",
        "Production output remains in the kitchen production/closing value ledger rather than fabricating store lots.",
        "Role-specific written SOP guides at /sops/<role>, with live links, role redirects, print-friendly layout, owner access to all guides, and language-cookie support.",
        "Reviewed Telugu title, timing, reason, and refusal copy for all role moments, plus global EN/Telugu toggle behavior.",
    ]),
    ("Sales, Petpooja, and POS reconciliation", [
        "Petpooja credential setup is owner-only, stores encrypted credentials, exposes only a non-secret configured indicator, and decrypts only at fetch time.",
        "Local demo POS mode works without real provider credentials.",
        "Petpooja Get Orders contract rejects successful-looking responses without order_json before normalization.",
        "Petpooja fetch retries network errors, HTTP 429, and HTTP 5xx responses up to three attempts; authentication and other client errors fail immediately.",
        "Durable POS sync-run records track running, succeeded, and failed attempts, attempt number, and visible errors.",
        "Latest-fetch-wins behavior preserves the newest normalized generation and reverses superseded journal postings exactly once.",
        "Sales CSV import validates IDs, dates, statuses, amounts, future dates, and duplicates before using the existing immutable POS persistence boundary.",
        "Sales CSV preview is write-free; editing the file clears the preview; commit reparses and revalidates.",
        "POS stock policy supports explicit none or reconcile behavior and never silently creates stock movements.",
        "POS statement CSV preview uses the same parser and validation as commit, with integer paise totals.",
        "POS reconciliation compares immutable provider statements against the latest revenue generation without overwriting orders or silently creating corrections.",
        "Difference review allows authorized sales/accounts users to acknowledge or request correction with a note while preserving the original statement and sales ledgers.",
        "Hourly cron endpoint uses a separate CRON_SECRET, gets tenant IDs through a narrow function, wraps every fetch in tenant context, and preserves sync-run/journal idempotency.",
    ]),
    ("Accounting and financial controls", [
        "Tenant-scoped chart of accounts, immutable journal headers/lines, balance enforcement, closed-period locking, and composite tenant ownership.",
        "Owner/accountant chart-of-accounts editor and accountant journal/trial-balance route.",
        "Explicit posting mappings; workflows refuse to post when required mappings are absent instead of guessing defaults.",
        "Vendor payments post payable liability and selected money-account asset in one transaction.",
        "Operating expenses post line-level journals and matching reversals on void.",
        "Other income and cash vouchers post using declared purpose and explicit asset or owner-payable mappings.",
        "Purchases post inventory, vendor payable, and explicitly configured input tax treatment atomically; voids reverse the exact original lines.",
        "Approved payroll payments post wages, net pay, advances, deductions, and withholding liabilities atomically.",
        "POS ingest posts revenue by explicit payment-mode asset mappings and sales-revenue account.",
        "Journal-backed P&L, cumulative balance sheet, direct cash flow, and indirect cash-flow bridge views.",
        "Recurring journal templates and idempotent period runs.",
        "Accrual headers, one-post-per-kind records, dated reversals, and balanced journal posting.",
        "Fixed-asset register and append-only straight-line depreciation postings with one-post-per-asset-period protection.",
        "Provider-neutral withholding CSV export with date range, bases, amounts, regime codes, derived rates, deposit/challan status, and notes; no filing or invented statutory rate.",
        "Opening-balance CSV import with preview, one-date scope, exact debit/credit balancing, active-account resolution, immutable batch evidence, and journal posting.",
    ]),
    ("Payroll, people, attendance, and leave", [
        "Effective-dated salary structures with positive-salary validation and history on Payroll → People.",
        "Payroll drafts select the latest structure effective by period end while preserving frozen prepared/paid runs.",
        "Printable payslips sourced from frozen payroll lines without recalculation or writes.",
        "Run-specific payroll CSV export using frozen payroll lines, formula-safe values, and explicit recorded-data semantics.",
        "Effective-dated statutory configuration for jurisdiction, PF/ESI percentage/cap, TDS regime, and source notes with immutable history.",
        "PF/ESI calculation helper using configured percentages and caps; no fabricated statutory defaults.",
        "Validated statutory input rejects invalid calendar dates and PF/ESI values above 100%.",
        "Tenant-scoped leave policies and effective staff assignments.",
        "Leave requests with overlap-safe pending/approved states, decision trail, manager/owner queue, cross-year refusal, and safe attendance integration.",
        "Approved leave writes one leave attendance fact per day without overwriting present or half-day marks.",
        "Formal carry-forward ledger with source-year policy bounds, approved-leave bounds, tenant locking, and no edit grant.",
        "Holiday calendar with explicit paid/unpaid policy; paid holidays apply only when no attendance mark exists.",
        "Validated staff CSV import with exact headers, employment/pay modes, amounts, dates, section references, duplicate checks, sequential employee codes, preview, and transaction-time revalidation.",
    ]),
    ("Evidence, attachments, and imports", [
        "Private attachment storage boundary supports local filesystem storage and the existing Vercel Blob backend.",
        "Authenticated streaming attachment reads with restaurant-first object keys.",
        "Attachment lifecycle metadata supports active/archived status and optional retention dates.",
        "Manager/owner archive controls are metadata-only; archived blobs remain available for audit and are not destructively deleted.",
        "Meter saves return immutable reading IDs and optionally attach evidence without delaying or rolling back the reading.",
        "Shared RFC-4180 CSV parser handles BOMs, quoted commas, embedded newlines, and malformed unclosed quotes consistently.",
        "Payroll, staff, vendor, item, opening stock, purchase, sales, POS, chart-account, and opening-balance imports all use preview-before-commit contracts where appropriate.",
        "Preview results report row counts, resolved references, proposed codes, totals, or approval mode without database writes.",
        "Commit actions reparse and revalidate inside the transaction so stale previews cannot bypass validation.",
    ]),
    ("User interface and operational reliability", [
        "Production landing page describing the complete KitchenBooks workflow and role responsibilities.",
        "Public root route while protected application pages still require authentication.",
        "Authenticated root route continues to redirect to the user’s role dashboard.",
        "Role-specific navigation and access-denied states.",
        "Build-boundary fixes for async Server Actions and server-only modules.",
        "Environment onboarding template separating required runtime values, temporary bootstrap PIN, local demo POS mode, and optional production capabilities.",
        "Production preflight fails closed for missing TLS, application secrets, private storage, malformed membership configuration, default PostgreSQL owner URLs, or demo POS mode in production.",
        "Migration inventory, migration-reference, migration-safety, deployment, schema, tenancy, role-matrix, and stock-lot audits are wired into the consolidated gates command.",
    ]),
]

for group, items in functional_groups:
    doc.add_heading(group, level=2)
    for item in items:
        add_bullet(doc, item)

doc.save(OUT)
print(OUT)
